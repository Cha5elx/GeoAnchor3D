from pathlib import Path
from copy import deepcopy
import re, hashlib, json
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
root=Path.cwd(); target=root/'docs/Response2Reviewers'; tmp=root/'tmp/word_reply'
template=Path('E:/个人资料/山东科技大学-硕士/多媒体智能计算-MIC/00 研究进展/审稿意见回复模板/TCSVT_Response Letter_R2.docx')
d=Document(template)
protos={k:deepcopy(d.paragraphs[i]._p.pPr) for k,i in [('heading',30),('commentlabel',31),('comment',60),('replylabel',61),('reply',62)]}
note='''# Word 回复信格式记录

格式来源：TCSVT_Response Letter_R2.docx，仅采用 Response to Reviewer 1、2、3。
- A4；上下页边距 2.54 cm，左右 3.175 cm。保留模板文档网格（312 twips）。
- Times New Roman；正文及 Comment/REPLY 标签 10.5 pt；审稿人标题 12 pt，加粗，采用模板 Heading 2 和原段落间距。
- Comment n: 独立一行、黑色、加粗、单下划线；意见正文黑色、左对齐。
- REPLY: 独立一行、蓝色 #3427E1、加粗、单下划线；回复正文同色，两端对齐，首行缩进 420 twips（约 0.741 cm）。
- 保留模板标题的自然连续排版；标签与后段保持同页，避免孤立标签。
- 模板中的页码/截图说明用橙色 #C45911。现有内容没有这类独立说明，不凭空增加。
- 仅格式迁移；英文 reviewer_1/2/3.tex 为内容依据。保留待补充及源文缺失标记，不导入模板论文、编者回复或结束语。
- 数学符号转换为可编辑 Word 上下标；不改写回复内容。
'''
(target/'word-response-template-style.md').write_text(note,encoding='utf-8')
(tmp/'artifact.md').write_text(note+'\nTemplate SHA256: '+hashlib.sha256(template.read_bytes()).hexdigest(),encoding='utf-8')
body=d._element.body
for x in list(body):
 if x.tag!=qn('w:sectPr'): body.remove(x)
# Drop template figures/comments and their unused relationships.
for rid,rel in list(d.part.rels.items()):
 if any(s in rel.reltype for s in ['/image','/comments','/people']): del d.part.rels[rid]
d.core_properties.title='GeoAnchor3D: Responses to Reviewers'
d.core_properties.author='';d.core_properties.last_modified_by='';d.core_properties.comments='';d.core_properties.subject=''

def run(p,text,kind,italic=False,bold=False,script=None):
 r=p.add_run(text); r.font.name='Times New Roman'; r.font.size=Pt(12 if kind=='heading' else 10.5)
 r.font.color.rgb=RGBColor.from_string('3427E1' if kind.startswith('reply') else '000000')
 r.bold=bold or kind in ['heading','commentlabel','replylabel'];r.italic=italic
 r.underline=kind in ['commentlabel','replylabel']
 if script=='_':r.font.subscript=True
 if script=='^':r.font.superscript=True
 return r

def group(s,i):
 assert s[i]=='{';depth=1;j=i+1
 while depth:
  if s[j]=='{':depth+=1
  if s[j]=='}':depth-=1
  j+=1
 return s[i+1:j-1],j

def inline(p,s,kind,italic=False,bold=False,script=None):
 i=0;buf=''
 def flush():
  nonlocal buf
  if buf:run(p,buf,kind,italic,bold,script);buf=''
 while i<len(s):
  ch=s[i]
  if ch=='\\':
   flush();m=re.match(r'\\([A-Za-z]+|.)',s[i:]);cmd=m.group(1);i+=len(m.group(0))
   if cmd in ['textit','textbf','mathbf','mathrm','mathcal','operatorname','url']:
    val,i=group(s,i);inline(p,val,kind,italic or cmd=='textit',bold or cmd in ['textbf','mathbf'],script)
   elif cmd in ['alpha','lambda','ast']:run(p,{'alpha':'α','lambda':'λ','ast':'∗'}[cmd],kind,italic,bold,script)
   elif cmd in ['%','&','_']:run(p,cmd,kind,italic,bold,script)
   elif cmd=='noindent':pass
   else:raise ValueError(cmd)
  elif ch in '_^':
   flush();i+=1
   if s[i]=='{':val,i=group(s,i)
   else:val=s[i];i+=1
   inline(p,val,kind,italic,bold,ch)
  elif ch=='$':i+=1
  elif ch=='~':buf+='\u00a0';i+=1
  elif s[i:i+3]=='---':buf+='—';i+=3
  elif s[i:i+2]=='--':buf+='–';i+=2
  elif s[i:i+2]=='``':buf+='“';i+=2
  elif s[i:i+2]=="''":buf+='”';i+=2
  else:buf+=ch;i+=1
 flush()

def para(text,kind):
 p=d.add_paragraph(); p._p.insert(0,deepcopy(protos[kind]))
 p.paragraph_format.widow_control=True
 if kind in ['heading','commentlabel','replylabel']:p.paragraph_format.keep_with_next=True
 inline(p,re.sub(r'\s+',' ',text).strip(),kind)
 return p
count=0;expected=[]
for reviewer in [1,2,3]:
 para(f'Response to Reviewer {reviewer}','heading')
 s=(target/f'reviewer_{reviewer}.tex').read_text(encoding='utf-8')
 s=re.sub(r'^%.*\n','',s,flags=re.M)
 pat=r'\{\\noindent\\color\{comment\}\s*\\textbf\{Comment (\d+):\}\s*(.*?)\}\s*\\vspace\{\\comafter\}(.*?)\\noindent\\textbf\{Response \1:\}\s*(.*?)\\vspace\{\\resafter\}'
 items=re.findall(pat,s,re.S)
 assert len(items)==(11 if reviewer==3 else 5),(reviewer,len(items))
 for num,comment,note,reply in items:
  para('Comment '+num+':','commentlabel');para(comment,'comment')
  if note.strip():para(note.strip(),'comment')
  para('REPLY:','replylabel')
  for part in re.split(r'\n\s*\n',reply.strip()):para(part,'reply')
  count+=1
  if int(num)!=len(items):d.add_paragraph()
assert count==21
out=target/'GeoAnchor3D_Response_Letter.docx';d.save(out)
plain='\n'.join(p.text for p in d.paragraphs)
assert plain.count('REPLY:')==21 and plain.count('Response to Reviewer ')==3
assert '\\' not in plain
(tmp/'content.txt').write_text(plain,encoding='utf-8')
print(str(out));print('Verified:',count,'comments and replies; characters:',len(plain))
