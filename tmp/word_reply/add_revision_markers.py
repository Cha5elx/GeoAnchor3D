from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
import re,json,hashlib
base=Path('docs/Response2Reviewers');src=base/'GeoAnchor3D_Response_Letter.docx'
d=Document(src)
# Each item: reviewer, comment, short introduction, source location, screenshot instructions.
items=[
(1,1,'The relevant formulation and the discussion of alternative spatial representations are reproduced below.','Sections III-B and V','图 A：III-B 从 “To ensure angular continuity” 至 “shared across all attention heads”，包含五维描述符、空间掩码公式及可学习矩阵说明（tex 198–212 行）。图 B：V 中从 “In addition, IGGA relies on handcrafted geometric cues” 到 “richer relational reasoning capabilities”（811 行中的两句）。前者是相关公式，并非声称全部为新增内容。'),
(1,2,'The revised discussion of proposal limitations is provided below.','Section V, Limitations and Future Work','截取 V 开头，从 “The performance of GeoAnchor3D still depends” 到 “dynamic proposal generation or open-vocabulary detection”（tex 811 行前半部分）。保留 missed、fragmented、merged 三种情况。'),
(1,3,'The efficiency comparison and the accompanying discussion are reproduced below.','Section IV-C, Computational Efficiency','图 A：完整效率表及表题 “Efficiency comparison on one NVIDIA RTX 3090”，包含训练和推理设置说明（tab:efficiency，tex 717–748 行）。图 B：Computational Efficiency 两段正文（750–753 行）。'),
(1,4,'The corresponding discussion of backbone transferability is reproduced below.','Section V, Limitations and Future Work','截取 V 中 “In GATH, the selection of intermediate supervision layers is currently predefined” 及紧随的 “Evaluating adaptive layer selection strategies with different language-model backbones…”（tex 811 行中部两句）。'),
(2,1,'The relevant IGGA formulations are reproduced below.','Section III-B, Instruction-Aware Geometric Gating Attention','按顺序截取相关公式及解释，可拆成三张：A 为逐头 Q/K/V 和语义注意力（tex 164–189 行）；B 为共享空间掩码、门控网络及维度说明（203–212、221–225 行）；C 为有效掩码、融合注意力和输出公式（234–250 行）。以论文当前内容为准，不额外补写 Normalize 定义或多头聚合说明。'),
(2,2,'The added analyses of gate supervision and within-task variation are provided below.','Section IV-C, Gating Granularity and Supervision; Within-Task Gate Variation','图 A：Gating Granularity and Supervision 中以 “Removing” 开头的段落（tex 543 行）。图 B：完整 “Within-task gate statistics on ScanRefer” 表及 Within-Task Gate Variation 两段（tab:within_task_gating，550–572 行）。表和段落若不相邻，可分开截图。'),
(2,3,'The controlled ablation and the clarification of the training settings are reproduced below.','Section IV-C, Ablation Study','图 A：IV-C 开头完整段落，说明两组消融使用的训练数据不同（tex 445 行）。图 B：完整 “Controlled ablation of gating granularity and gate supervision” 表，连同表题和设置说明（tab:gating_granularity_prior，467–510 行）。'),
(2,4,'The paragraph clarifying the inherited framework and the proposed modules is reproduced below.','Section III-A, Overview of GeoAnchor3D','截取从 “GeoAnchor3D follows the object-centric scene representation” 到 “the language-model backbone” 的完整蓝色段落（tex 147 行）。'),
(2,5,'The revised statements defining the evaluation scope are provided below.','Sections V and VI','图 A：V 中从 “All evaluated benchmarks are derived from ScanNet indoor scenes” 到该节结尾（tex 811 行后半部分）。图 B：Conclusion 最后两句，从 “Experiments on four ScanNet-based benchmarks” 开始（815 行）。'),
(3,1,'The added gate-prior ablation and within-task analysis are reproduced below.','Section IV-C, Gating Granularity and Supervision; Within-Task Gate Variation','图 A：完整受控门控消融表，包含 w/o gate prior 和 Full per-head 行（tab:gating_granularity_prior，tex 467–510 行）。图 B：完整任务内门控统计表及其两段分析（tab:within_task_gating，550–572 行）。可复用 R2-C3 和 R2-C2 对应截图。'),
(3,2,'The scalar-versus-per-head comparison and its training setting are provided below.','Section IV-C, Ablation Study','图 A：IV-C 开头关于两组训练数据设置的段落（tex 445 行）。图 B：完整受控门控消融表（tab:gating_granularity_prior，467–510 行）。可复用 R2-C3 截图；不安排尚未开展的 layer-wise gating 实验截图。'),
(3,3,'The layer-wise probing results and the separate component ablation are reproduced below.','Section IV-C, Layer-Wise Geometric Recoverability; Module Contributions','图 A：完整 “Layer-wise linear probing of scene-centered object coordinates” 表及 Layer-Wise Geometric Recoverability 两段，保留冻结模型、场景划分及完整框架证据范围（tab:layerwise_probe，tex 674–711 行）。图 B：完整 IGGA/GATH 组件消融表（tab:final_ablation，418–439 行），用于支持 GATH 的任务贡献。'),
(3,4,'The revised description of the layer-interval results is reproduced below.','Section IV-C, Feature Extraction Depth','截取 Feature Extraction Depth 中从 “Beyond the gating design” 到 “performance varies across the selected layer intervals” 的前三句，包含改后的蓝色句子（tex 598 行）。表中数值已在回复中列出，此处不必重复贴整张层区间表。'),
(3,6,'The revised statement on unseen-task generalization is reproduced below.','Section V, Limitations and Future Work','截取 V 中从 “All evaluated benchmarks are derived from ScanNet indoor scenes” 到该节结尾（tex 811 行后半部分），包含训练和评估任务重合、within-domain multi-task adaptation 及 held-out task future work。'),
(3,7,'The dataset-domain clarification and the revised conclusion are provided below.','Section IV-A, Datasets; Section VI, Conclusion','图 A：Datasets 中从 “Although the four benchmarks cover different task formulations” 到 “indoor-scene domain” 的两句蓝字（tex 383 行）。图 B：Conclusion 最后两句蓝字（815 行）。可复用 R2-C5 的结论截图。'),
(3,8,'The efficiency measurements and the clarification of inference-time costs are reproduced below.','Section IV-C, Computational Efficiency','图 A：完整效率表及表题中的测量设置（tab:efficiency，tex 717–748 行）。图 B：Computational Efficiency 两段，特别保留 “GATH … are not executed, while IGGA remains active” 及相同输入和配置说明（750–753 行）。可复用 R1-C3 截图。'),
(3,9,'The clarification of the inherited proposals and their limitations is reproduced below.','Section III-A, Overview of GeoAnchor3D; Section V, Limitations and Future Work','图 A：III-A 从 “GeoAnchor3D follows” 到 “the language-model backbone” 的完整段落（tex 147 行，可复用 R2-C4）。图 B：V 开头至 “dynamic proposal generation or open-vocabulary detection”（811 行，可复用 R1-C2）。'),
(3,10,'The revised t-SNE discussion and the quantitative coordinate-probing results are provided below.','Section IV-D, Training Dynamics and Feature Geometry; Section IV-C, Layer-Wise Geometric Recoverability','图 A：IV-D 从 “To visualize the geometry of the intermediate representations” 开始的完整段落（tex 806 行），包含定性聚类与坐标可恢复性结果的联系。图 B：完整坐标探测表及结果段落（tab:layerwise_probe，674–711 行，可复用 R3-C3 图 A）。无需重复贴原有 t-SNE 图，重点展示修改后的解释。'),
(3,11,'The added discussion and the corresponding references are reproduced below.','Section II-A, Cross-Modal Scene Understanding; References','图 A：II-A 中 “The importance of preserving spatial structure has also been recognized…” 完整蓝色句子及三个引用编号（tex 122 行）。图 B：最终编译 PDF 的 References 中对应的三条完整文献：huang2026underwater、gao2026arctic、huang2026contourlet。引用编号随最终编译结果确定。')
]
lookup={(r,c):(lead,loc,note) for r,c,lead,loc,note in items}
# Find original response ends without disturbing existing run formatting.
blocks={};reviewer=0;comment=0;inreply=False
for p in d.paragraphs:
 t=p.text.strip()
 if t.startswith('Response to Reviewer '):reviewer=int(t[-1]);inreply=False
 elif re.fullmatch(r'Comment \d+:',t):comment=int(re.search(r'\d+',t).group());inreply=False
 elif t=='REPLY:':inreply=True
 elif inreply and t:blocks.setdefault((reviewer,comment),[]).append(p)
assert len(blocks)==21
removed=[]
for key,paras in blocks.items():
 for p in list(paras):
  if p.text.startswith('[Manuscript alignment pending:'):
   removed.append(p.text);p._p.getparent().remove(p._p);paras.remove(p)

def insert(after,text,color='3427E1',italic=False,keep=False):
 e=OxmlElement('w:p');after._p.addnext(e);p=Paragraph(e,after._parent)
 p._p.insert(0,deepcopy(after._p.pPr))
 p.paragraph_format.keep_with_next=keep;p.paragraph_format.widow_control=True
 p.paragraph_format.first_line_indent=Pt(0)
 r=p.add_run(text);r.font.name='Times New Roman';r.font.size=Pt(10.5);r.font.color.rgb=RGBColor.from_string(color);r.italic=italic;r.bold=False;r.underline=False
 r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
 return p
for key,(lead,loc,note) in lookup.items():
 end=blocks[key][-1]
 p=insert(end,lead+' ['+loc+']',keep=True)
 # Location follows template's orange underlined source references.
 p.runs[0].text=lead+' '
 r=p.add_run('['+loc+']');r.font.name='Times New Roman';r.font.size=Pt(10.5);r.font.color.rgb=RGBColor.from_string('C45911');r.underline=True
 insert(p,f'【截图位置 R{key[0]}-C{key[1]}：{note} 插图后删除本标注。】','C45911')
out=base/'GeoAnchor3D_Response_Letter_with_revision_markers.docx';d.save(out)
check=Document(out)
assert sum(p.text=='REPLY:' for p in check.paragraphs)==21
assert sum(p.text.startswith('【截图位置') for p in check.paragraphs)==19
# All original paragraphs except the resolved editorial note must remain verbatim.
old=[p.text for p in Document(src).paragraphs if p.text not in removed]
new=[p.text for p in check.paragraphs]
it=iter(new)
assert all(any(t==v for v in it) for t in old)
paper='E:/个人资料/山东科技大学-硕士/多媒体智能计算-MIC/00 研究进展/02 工作二 ScaneGraphReasoning/SGR/docs/GeoAnchor3D.tex'
notes=['# 回复信截图清单','',f'论文依据：{paper}','','模板句式：Below is our revised content [章节/页码/段落]:，随后插入修改后的论文截图。此版按语境使用 revised discussion、added analyses 或 relevant formulation，避免把未改动的公式统称新增内容。','','19 处截图位置已放入 Word；橙色中文仅供编辑，截图后删除。英文位置说明保留，页码可在最终 PDF 确定后补上。应截编译后的论文 PDF，不截 LaTeX 编辑器；保留表题、列名、公式编号及修改颜色，避免截整页。每张图按正文宽度等比插入，宽表横跨正文宽度，过长内容分图，保证清晰。','','R1-C5、R3-C5：实验结果仍待补充，本版不新增引导句或截图占位。R3-C1 的原始审稿意见缺失符号提示仍保留。R3-C10 的 t-SNE 待同步提示已据 tex 806 行落实的内容替换成截图引导。']
for r,c,lead,loc,note in items:notes += ['',f'## R{r}-C{c}',lead+' ['+loc+']',note]
(base/'回复信截图清单.md').write_text('\n'.join(notes),encoding='utf-8')
Path('tmp/word_reply/revision-marker-audit.json').write_text(json.dumps({'original_sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'added_locations':19,'unchanged_comments_and_responses':21,'resolved_note':removed,'output':str(out)},ensure_ascii=False,indent=2),encoding='utf-8')
print('PASS: 19 screenshot locations; 21 responses preserved; resolved 1 obsolete editorial note.')
